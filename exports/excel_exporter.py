import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def export_to_excel(rows, headers, filename="export.xlsx") -> str:
    """
    Genera un archivo Excel con los datos proporcionados,
    lo guarda y lo abre automáticamente con el programa predeterminado.
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # ── Estilo de encabezado ──
    header_fill = PatternFill("solid", fgColor="1A1D27")
    header_font = Font(bold=True, color="3DD68C", size=11)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill   = header_fill
        cell.font   = header_font
        cell.alignment = header_align

    # ── Datos ──
    row_font_even = Font(size=10)
    row_font_odd  = Font(size=10)
    fill_even     = PatternFill("solid", fgColor="F2F2F2")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if row_idx % 2 == 0:
                cell.fill = fill_even
                cell.font = row_font_even
            else:
                cell.font = row_font_odd

    # ── Ajustar ancho de columnas ──
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    # ── Guardar ──
    wb.save(filename)

    # ── Abrir automáticamente ──
    _open_file(filename)

    return os.path.abspath(filename)


def _open_file(path: str):
    """Abre el archivo con la aplicación predeterminada del sistema."""
    try:
        os.startfile(path)          # Windows
    except AttributeError:
        import subprocess
        try:
            subprocess.run(["xdg-open", path], check=False)   # Linux
        except Exception:
            subprocess.run(["open", path], check=False)        # macOS
    except Exception:
        pass  # Si falla, no bloquea el flujo
